"""
Styled Excel export for the AVL-DRIVE HeatMap view.

Builds a workbook that mirrors the on-screen heatmap: header rows, separator
columns, score color-scale fills, status dots, and parent group labels.
"""

from __future__ import annotations

import io
from typing import Optional

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from config import (
    COLOR_BLACK,
    COLOR_GREEN,
    COLOR_HEATMAP_BORDER,
    COLOR_HEATMAP_HEADER,
    COLOR_RED,
    COLOR_WHITE,
    COLOR_YELLOW,
    COLOR_YELLOW_BRIGHT,
    DOT_FONT_COLORS,
    PARENT_OPERATION_CODES,
    SCORE_COLOR_MAX,
    SCORE_COLOR_MID,
    SCORE_COLOR_MIN,
    SCORE_SCALE_MAX,
    SCORE_SCALE_MID,
    SCORE_SCALE_MIN,
)

# Status dot colors match the on-screen HeatMap HTML renderer.
_STATUS_DOT_COLORS = {
    "GREEN": COLOR_GREEN,
    "YELLOW": COLOR_YELLOW,
    "RED": COLOR_RED,
    "BLUE": DOT_FONT_COLORS["BLUE"],
}

_THIN = Side(style="thin", color=COLOR_HEATMAP_BORDER.lstrip("#"))
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
_NO_BORDER = Border()
_SEP_BORDER = Border(top=_THIN, bottom=_THIN)

_FONT_DEFAULT = Font(name="Calibri", size=11, color=COLOR_BLACK.lstrip("#"))
_FONT_BOLD = Font(name="Calibri", size=11, bold=True, color=COLOR_BLACK.lstrip("#"))
_FONT_HEADER = Font(name="Calibri", size=11, bold=True, color=COLOR_BLACK.lstrip("#"))
_FONT_TARGET = Font(name="Calibri", size=10, bold=True, color=COLOR_BLACK.lstrip("#"))
_FONT_DOT = Font(name="Calibri", size=16, bold=True)
_FONT_STATUS_PARENT = Font(name="Calibri", size=11, bold=True)

_ALIGN_CENTER = Alignment(horizontal="center", vertical="center")
_ALIGN_LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)
_ALIGN_LEFT_NOWRAP = Alignment(horizontal="left", vertical="center")

_FILL_HEADER = PatternFill("solid", fgColor=COLOR_HEATMAP_HEADER.lstrip("#"))
_FILL_WHITE = PatternFill("solid", fgColor=COLOR_WHITE.lstrip("#"))


def export_heatmap_to_excel(
    df: pd.DataFrame,
    target_vehicle: Optional[str] = None,
    tested_vehicle: Optional[str] = None,
    sheet_name: str = "HeatMap",
) -> io.BytesIO:
    """Export a styled heatmap DataFrame to an Excel workbook buffer."""
    vehicle_cols = _vehicle_columns(df)
    vehicle_cols = _order_vehicle_columns(vehicle_cols, target_vehicle, tested_vehicle)
    has_status = "Status" in df.columns
    has_comments = "Comments" in df.columns

    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name

    layout = _build_column_layout(vehicle_cols, has_status, has_comments)
    _apply_column_widths(ws, layout)

    row = 1
    _write_target_tested_row(
        ws, row, layout, target_vehicle=target_vehicle, tested_vehicle=tested_vehicle
    )
    row += 1
    _write_header_row(ws, row, layout, vehicle_cols, has_status, has_comments)
    row += 1
    _write_dr_row(ws, row, layout, has_status, has_comments)
    row += 1

    for _, data_row in df.iterrows():
        _write_data_row(ws, row, layout, data_row, has_status, has_comments)
        row += 1

    ws.freeze_panes = "C4"
    ws.sheet_view.showGridLines = False

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def _vehicle_columns(df: pd.DataFrame) -> list[str]:
    return [
        c
        for c in df.columns
        if c not in ("Op Code", "Operation Mode", "Status", "Comments")
    ]


def _order_vehicle_columns(
    vehicle_cols: list[str],
    target_vehicle: Optional[str],
    tested_vehicle: Optional[str],
) -> list[str]:
    if not target_vehicle or not tested_vehicle or target_vehicle == tested_vehicle:
        return list(vehicle_cols)

    middle = [c for c in vehicle_cols if c not in (target_vehicle, tested_vehicle)]
    ordered: list[str] = []
    if target_vehicle in vehicle_cols:
        ordered.append(target_vehicle)
    ordered.extend(middle)
    if tested_vehicle in vehicle_cols:
        ordered.append(tested_vehicle)
    return ordered


def _build_column_layout(
    vehicle_cols: list[str],
    has_status: bool,
    has_comments: bool,
) -> list[tuple[str, Optional[str]]]:
    layout: list[tuple[str, Optional[str]]] = [
        ("op_code", None),
        ("op_mode", None),
    ]
    for vname in vehicle_cols:
        layout.append(("sep", vname))
        layout.append(("score", vname))
    if has_status:
        layout.append(("status", None))
    if has_comments:
        layout.append(("comments", None))
    return layout


def _apply_column_widths(ws, layout: list[tuple[str, Optional[str]]]) -> None:
    widths = {
        "op_code": 11,
        "op_mode": 34,
        "sep": 1.2,
        "score": 14,
        "status": 11,
        "comments": 42,
    }
    for col_idx, (kind, _) in enumerate(layout, start=1):
        letter = get_column_letter(col_idx)
        ws.column_dimensions[letter].width = widths[kind]


def _write_target_tested_row(
    ws,
    row: int,
    layout: list[tuple[str, Optional[str]]],
    *,
    target_vehicle: Optional[str],
    tested_vehicle: Optional[str],
) -> None:
    for col_idx, (kind, vname) in enumerate(layout, start=1):
        cell = ws.cell(row=row, column=col_idx)
        if kind == "op_code":
            cell.value = ""
            _style_cell(cell, fill=_FILL_WHITE, border=_NO_BORDER)
        elif kind == "op_mode":
            cell.value = ""
            _style_cell(cell, fill=_FILL_WHITE, border=_NO_BORDER)
        elif kind == "sep":
            cell.value = ""
            _style_cell(cell, fill=_FILL_WHITE, border=_NO_BORDER)
        elif kind == "score":
            label = ""
            if vname == target_vehicle:
                label = "Target Vehicle"
            elif vname == tested_vehicle:
                label = "Tested Vehicle"
            cell.value = label
            _style_cell(
                cell,
                fill=_FILL_WHITE,
                font=_FONT_TARGET,
                alignment=_ALIGN_CENTER,
                border=_BORDER if label else _NO_BORDER,
            )
        else:
            cell.value = ""
            _style_cell(cell, fill=_FILL_WHITE, border=_NO_BORDER)


def _write_header_row(
    ws,
    row: int,
    layout: list[tuple[str, Optional[str]]],
    vehicle_cols: list[str],
    has_status: bool,
    has_comments: bool,
) -> None:
    for col_idx, (kind, vname) in enumerate(layout, start=1):
        cell = ws.cell(row=row, column=col_idx)
        if kind == "op_code":
            cell.value = ""
            _style_cell(cell, fill=_FILL_HEADER, font=_FONT_HEADER, border=_BORDER)
        elif kind == "op_mode":
            cell.value = "Operation Modes"
            _style_cell(
                cell,
                fill=_FILL_HEADER,
                font=_FONT_HEADER,
                alignment=_ALIGN_LEFT_NOWRAP,
                border=_BORDER,
            )
        elif kind == "sep":
            cell.value = ""
            _style_cell(cell, fill=_FILL_WHITE, border=_SEP_BORDER)
        elif kind == "score":
            cell.value = vname
            _style_cell(
                cell,
                fill=_FILL_HEADER,
                font=_FONT_HEADER,
                alignment=_ALIGN_CENTER,
                border=_BORDER,
            )
        elif kind == "status" and has_status:
            cell.value = "Status"
            _style_cell(
                cell,
                fill=_FILL_HEADER,
                font=_FONT_HEADER,
                alignment=_ALIGN_CENTER,
                border=_BORDER,
            )
        elif kind == "comments" and has_comments:
            cell.value = "Comments"
            _style_cell(
                cell,
                fill=_FILL_HEADER,
                font=_FONT_HEADER,
                alignment=_ALIGN_CENTER,
                border=_BORDER,
            )


def _write_dr_row(
    ws,
    row: int,
    layout: list[tuple[str, Optional[str]]],
    has_status: bool,
    has_comments: bool,
) -> None:
    for col_idx, (kind, _) in enumerate(layout, start=1):
        cell = ws.cell(row=row, column=col_idx)
        if kind == "op_code":
            cell.value = ""
            _style_cell(cell, fill=_FILL_HEADER, border=_BORDER)
        elif kind == "op_mode":
            cell.value = ""
            _style_cell(cell, fill=_FILL_HEADER, border=_BORDER)
        elif kind == "sep":
            cell.value = ""
            _style_cell(cell, fill=_FILL_WHITE, border=_SEP_BORDER)
        elif kind == "score":
            cell.value = "DR"
            _style_cell(
                cell,
                fill=_FILL_HEADER,
                font=_FONT_DEFAULT,
                alignment=_ALIGN_CENTER,
                border=_BORDER,
            )
        elif kind == "status" and has_status:
            cell.value = ""
            _style_cell(cell, fill=_FILL_HEADER, border=_BORDER)
        elif kind == "comments" and has_comments:
            cell.value = ""
            _style_cell(cell, fill=_FILL_HEADER, border=_BORDER)


def _write_data_row(
    ws,
    row: int,
    layout: list[tuple[str, Optional[str]]],
    data_row: pd.Series,
    has_status: bool,
    has_comments: bool,
) -> None:
    op_code = data_row.get("Op Code", "")
    op_name = data_row.get("Operation Mode", "")
    is_parent = op_code in PARENT_OPERATION_CODES
    row_fill = _FILL_HEADER if is_parent else _FILL_WHITE
    row_font = _FONT_BOLD if is_parent else _FONT_DEFAULT

    status_val = data_row.get("Status", "") if has_status else ""
    status_str = _clean_str(status_val)
    status_upper = status_str.upper()

    for col_idx, (kind, vname) in enumerate(layout, start=1):
        cell = ws.cell(row=row, column=col_idx)

        if kind == "op_code":
            cell.value = op_code if op_code != "" else None
            _style_cell(
                cell,
                fill=row_fill,
                font=row_font,
                alignment=_ALIGN_CENTER,
                border=_BORDER,
            )
        elif kind == "op_mode":
            cell.value = str(op_name) if op_name is not None else ""
            _style_cell(
                cell,
                fill=row_fill,
                font=row_font,
                alignment=_ALIGN_LEFT_NOWRAP,
                border=_BORDER,
            )
        elif kind == "sep":
            cell.value = ""
            _style_cell(cell, fill=_FILL_WHITE, border=_SEP_BORDER)
        elif kind == "score":
            val = data_row.get(vname)
            cell.value = _score_value(val)
            bg_hex, _ = _score_bg(val)
            _style_cell(
                cell,
                fill=_solid_fill(bg_hex),
                font=_FONT_DEFAULT,
                alignment=_ALIGN_CENTER,
                border=_BORDER,
            )
        elif kind == "status" and has_status:
            _write_status_cell(cell, is_parent, status_str, status_upper)
        elif kind == "comments" and has_comments:
            comment_val = data_row.get("Comments", "")
            cell.value = _clean_str(comment_val)
            _style_cell(
                cell,
                fill=_FILL_WHITE,
                font=_FONT_DEFAULT,
                alignment=_ALIGN_LEFT,
                border=_BORDER,
            )


def _write_status_cell(cell, is_parent: bool, status_str: str, status_upper: str) -> None:
    if is_parent:
        if status_upper == "NOK":
            fill, font_color = COLOR_RED, COLOR_WHITE
        elif status_upper == "ACCEPTABLE":
            fill, font_color = COLOR_YELLOW_BRIGHT, COLOR_BLACK
        elif status_upper == "OK":
            fill, font_color = COLOR_GREEN, COLOR_WHITE
        else:
            fill, font_color = COLOR_WHITE, COLOR_BLACK
        cell.value = status_str
        _style_cell(
            cell,
            fill=_solid_fill(fill),
            font=Font(
                name="Calibri",
                size=11,
                bold=True,
                color=font_color.lstrip("#"),
            ),
            alignment=_ALIGN_CENTER,
            border=_BORDER,
        )
        return

    cell.value = "●" if status_upper in _STATUS_DOT_COLORS else ""
    if status_upper in _STATUS_DOT_COLORS:
        dot_color = _STATUS_DOT_COLORS[status_upper].lstrip("#")
        _style_cell(
            cell,
            fill=_FILL_WHITE,
            font=Font(name="Calibri", size=16, bold=True, color=dot_color),
            alignment=_ALIGN_CENTER,
            border=_BORDER,
        )
    else:
        _style_cell(
            cell,
            fill=_FILL_WHITE,
            font=_FONT_DEFAULT,
            alignment=_ALIGN_CENTER,
            border=_BORDER,
        )


def _style_cell(
    cell,
    *,
    fill: Optional[PatternFill] = None,
    font: Optional[Font] = None,
    alignment: Optional[Alignment] = None,
    border: Optional[Border] = None,
) -> None:
    if fill is not None:
        cell.fill = fill
    if font is not None:
        cell.font = font
    if alignment is not None:
        cell.alignment = alignment
    if border is not None:
        cell.border = border


def _solid_fill(hex_color: str) -> PatternFill:
    return PatternFill("solid", fgColor=hex_color.lstrip("#"))


def _clean_str(val) -> str:
    if val is None or (isinstance(val, float) and pd.isna(val)):
        return ""
    return str(val).strip()


def _score_value(val):
    text = _fmt_score(val)
    if not text:
        return None
    try:
        if "." in text:
            return float(text)
        return int(text)
    except ValueError:
        return text


def _fmt_score(val) -> str:
    if val is None or (isinstance(val, float) and pd.isna(val)):
        return ""
    try:
        v = float(val)
    except (ValueError, TypeError):
        return str(val)
    if v <= 0:
        return ""
    return str(int(v)) if v == int(v) else f"{v:.1f}"


def _score_bg(val) -> tuple[str, str]:
    if val is None or (isinstance(val, float) and pd.isna(val)):
        return COLOR_WHITE, COLOR_BLACK
    try:
        v = float(val)
    except (ValueError, TypeError):
        return COLOR_WHITE, COLOR_BLACK
    if v <= 0:
        return COLOR_WHITE, COLOR_BLACK
    return _score_gradient(v), COLOR_BLACK


def _score_gradient(v: float) -> str:
    if v <= SCORE_SCALE_MIN:
        r, g, b = SCORE_COLOR_MIN
    elif v >= SCORE_SCALE_MAX:
        r, g, b = SCORE_COLOR_MAX
    elif v <= SCORE_SCALE_MID:
        t = (v - SCORE_SCALE_MIN) / (SCORE_SCALE_MID - SCORE_SCALE_MIN)
        r = int(SCORE_COLOR_MIN[0] + t * (SCORE_COLOR_MID[0] - SCORE_COLOR_MIN[0]))
        g = int(SCORE_COLOR_MIN[1] + t * (SCORE_COLOR_MID[1] - SCORE_COLOR_MIN[1]))
        b = int(SCORE_COLOR_MIN[2] + t * (SCORE_COLOR_MID[2] - SCORE_COLOR_MIN[2]))
    else:
        t = (v - SCORE_SCALE_MID) / (SCORE_SCALE_MAX - SCORE_SCALE_MID)
        r = int(SCORE_COLOR_MID[0] + t * (SCORE_COLOR_MAX[0] - SCORE_COLOR_MID[0]))
        g = int(SCORE_COLOR_MID[1] + t * (SCORE_COLOR_MAX[1] - SCORE_COLOR_MID[1]))
        b = int(SCORE_COLOR_MID[2] + t * (SCORE_COLOR_MAX[2] - SCORE_COLOR_MID[2]))
    return f"#{r:02X}{g:02X}{b:02X}"

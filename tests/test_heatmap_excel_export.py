"""Tests for styled HeatMap Excel export."""

import io

import pandas as pd
import pytest
from openpyxl import load_workbook

from heatmap_excel_export import export_heatmap_to_excel


def _sample_heatmap_df() -> pd.DataFrame:
    return pd.DataFrame(
        [
            {
                "Op Code": 10100000,
                "Operation Mode": "Drive away",
                "Target Car": 7.9,
                "Tested Car": 7.4,
                "Status": "OK",
            },
            {
                "Op Code": 10101300,
                "Operation Mode": "Creep",
                "Target Car": 7.2,
                "Tested Car": 6.9,
                "Status": "RED",
                "Comments": "Red P1 Drivability, Brake Release Bump",
            },
            {
                "Op Code": 10120000,
                "Operation Mode": "Acceleration",
                "Target Car": 8.1,
                "Tested Car": 8.0,
                "Status": "Acceptable",
            },
        ]
    )


def _load_exported_workbook(df, **kwargs):
    buf = export_heatmap_to_excel(df, **kwargs)
    return load_workbook(io.BytesIO(buf.getvalue()))


def test_export_creates_expected_sheet_layout():
    wb = _load_exported_workbook(
        _sample_heatmap_df(),
        target_vehicle="Target Car",
        tested_vehicle="Tested Car",
    )
    ws = wb["HeatMap"]

    assert ws["B2"].value == "Operation Modes"
    assert ws["D2"].value == "Target Car"
    assert ws["F2"].value == "Tested Car"
    assert ws["D3"].value == "DR"
    assert ws["F3"].value == "DR"
    assert ws["D1"].value == "Target Vehicle"
    assert ws["F1"].value == "Tested Vehicle"
    assert ws["G2"].value == "Status"
    assert ws["H2"].value == "Comments"


def test_export_applies_score_color_fills():
    wb = _load_exported_workbook(
        _sample_heatmap_df(),
        target_vehicle="Target Car",
        tested_vehicle="Tested Car",
    )
    ws = wb["HeatMap"]

    # Row 4 = first data row (Drive away), Target Car score in column D.
    score_cell = ws["D4"]
    assert score_cell.value == 7.9
    high_score_rgb = score_cell.fill.fgColor.rgb
    assert high_score_rgb not in (None, "00000000", "00FFFFFF", "FFFFFFFF")

    # Lower score should have a warmer (more red/yellow) fill than the higher score.
    creep_score = ws["F5"]
    assert creep_score.value == 6.9
    low_score_rgb = creep_score.fill.fgColor.rgb
    assert low_score_rgb not in (None, "00000000", "00FFFFFF", "FFFFFFFF")
    assert low_score_rgb != high_score_rgb


def test_export_renders_status_dots_and_parent_labels():
    wb = _load_exported_workbook(
        _sample_heatmap_df(),
        target_vehicle="Target Car",
        tested_vehicle="Tested Car",
    )
    ws = wb["HeatMap"]

    parent_status = ws["G4"]
    assert parent_status.value == "OK"
    assert parent_status.font.bold is True
    assert parent_status.fill.start_color.rgb in ("0000B050", "00B050")

    sub_status = ws["G5"]
    assert sub_status.value == "●"
    assert sub_status.font.color.rgb in ("00FF0000", "FFFF0000")

    acceptable_status = ws["G6"]
    assert acceptable_status.value == "Acceptable"
    assert acceptable_status.fill.start_color.rgb in ("00FFFF00", "FFFF00")


def test_export_includes_comments():
    wb = _load_exported_workbook(
        _sample_heatmap_df(),
        target_vehicle="Target Car",
        tested_vehicle="Tested Car",
    )
    ws = wb["HeatMap"]

    assert "Brake Release Bump" in ws["H5"].value


def test_export_without_status_columns():
    df = _sample_heatmap_df().drop(columns=["Status", "Comments"])
    wb = _load_exported_workbook(
        df,
        target_vehicle="Target Car",
        tested_vehicle="Tested Car",
    )
    ws = wb["HeatMap"]

    assert ws.max_column == 6  # op code, op mode, sep+score x2
    assert ws["B2"].value == "Operation Modes"
    assert ws["D4"].value == 7.9


def test_score_gradient_endpoints():
    df = pd.DataFrame(
        [
            {
                "Op Code": 10101300,
                "Operation Mode": "Creep",
                "Target Car": 1.0,
                "Tested Car": 10.0,
            },
        ]
    )
    wb = _load_exported_workbook(
        df,
        target_vehicle="Target Car",
        tested_vehicle="Tested Car",
    )
    ws = wb["HeatMap"]

    low_rgb = ws["D4"].fill.fgColor.rgb
    high_rgb = ws["F4"].fill.fgColor.rgb
    assert low_rgb in ("00FF0000", "FFFF0000")
    assert high_rgb in ("0000B050", "00B050")


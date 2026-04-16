"""
Unit tests for the AVL-DRIVE Heatmap Tool evaluation logic.
"""

import sys
import os
import pandas as pd

# Add project root to path
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from evaluation_engine import (
    _evaluate_status,
    _combine_status,
    _bench_diff,
    _to_float,
    _parse_dot_status,
    evaluate_avl_status,
    build_overall_status,
    calculate_group_status,
    update_sub_operation_heatmap,
    parse_sheet1_data,
    parse_sheet1_from_excel,
)
from heatmap_engine import (
    build_heatmap_template,
    parse_heatmap_data,
    refresh_heatmap,
    filter_heatmap_rows,
)
from config import BENCH_SENTINEL, PARENT_OPERATION_CODES


# ============================================================================
# Tests for _evaluate_status
# ============================================================================
class TestEvaluateStatus:
    def test_p1_na_returns_na(self):
        assert _evaluate_status(8.0, "N/A", 0, 100, 100) == "N/A"

    def test_avl_below_7_returns_red(self):
        assert _evaluate_status(6.5, "GREEN", 0, 100, 100) == "RED"

    def test_p1_red_returns_red(self):
        assert _evaluate_status(8.0, "RED", 0, 100, 100) == "RED"

    def test_p1_yellow_returns_yellow(self):
        assert _evaluate_status(8.0, "YELLOW", 0, 100, 100) == "YELLOW"

    def test_green_no_benchmark_returns_green(self):
        assert _evaluate_status(8.0, "GREEN", BENCH_SENTINEL, 0, 0) == "GREEN"

    def test_green_tested_ge_target_returns_green(self):
        assert _evaluate_status(8.0, "GREEN", 0, 90, 95) == "GREEN"

    def test_green_within_tolerance_returns_green(self):
        assert _evaluate_status(8.0, "GREEN", 1.5, 95, 93.5) == "GREEN"

    def test_green_exceeds_tolerance_returns_yellow(self):
        assert _evaluate_status(8.0, "GREEN", 5.0, 100, 95) == "YELLOW"

    def test_avl_exactly_7_with_green_p1(self):
        assert _evaluate_status(7.0, "GREEN", BENCH_SENTINEL, 0, 0) == "GREEN"

    def test_avl_exactly_7_with_yellow_p1(self):
        assert _evaluate_status(7.0, "YELLOW", 0, 100, 100) == "YELLOW"


# ============================================================================
# Tests for _combine_status
# ============================================================================
class TestCombineStatus:
    def test_both_green(self):
        assert _combine_status("GREEN", "GREEN") == "GREEN"

    def test_either_red(self):
        assert _combine_status("RED", "GREEN") == "RED"
        assert _combine_status("GREEN", "RED") == "RED"

    def test_either_yellow(self):
        assert _combine_status("YELLOW", "GREEN") == "YELLOW"
        assert _combine_status("GREEN", "YELLOW") == "YELLOW"

    def test_green_and_na(self):
        assert _combine_status("GREEN", "N/A") == "GREEN"
        assert _combine_status("N/A", "GREEN") == "GREEN"

    def test_both_na(self):
        assert _combine_status("N/A", "N/A") == "N/A"

    def test_red_over_yellow(self):
        assert _combine_status("RED", "YELLOW") == "RED"

    def test_yellow_and_na(self):
        assert _combine_status("YELLOW", "N/A") == "YELLOW"


# ============================================================================
# Tests for _bench_diff
# ============================================================================
class TestBenchDiff:
    def test_both_zero(self):
        assert _bench_diff(0, 0) == BENCH_SENTINEL

    def test_target_zero(self):
        assert _bench_diff(0, 50) == BENCH_SENTINEL

    def test_normal_diff(self):
        assert _bench_diff(100, 95) == 5.0

    def test_tested_higher(self):
        assert _bench_diff(90, 100) == 10.0


# ============================================================================
# Tests for _to_float
# ============================================================================
class TestToFloat:
    def test_valid_number(self):
        assert _to_float("8.3") == 8.3

    def test_int_string(self):
        assert _to_float("100") == 100.0

    def test_empty_string(self):
        assert _to_float("") == 0.0

    def test_none(self):
        assert _to_float(None) == 0.0

    def test_non_numeric(self):
        assert _to_float("abc") == 0.0


# ============================================================================
# Tests for _parse_dot_status
# ============================================================================
class TestParseDotStatus:
    def test_green(self):
        assert _parse_dot_status("G") == "GREEN"
        assert _parse_dot_status("GREEN") == "GREEN"

    def test_yellow(self):
        assert _parse_dot_status("Y") == "YELLOW"
        assert _parse_dot_status("YELLOW") == "YELLOW"

    def test_red(self):
        assert _parse_dot_status("R") == "RED"
        assert _parse_dot_status("RED") == "RED"

    def test_blank(self):
        assert _parse_dot_status("") == "N/A"
        assert _parse_dot_status(None) == "N/A"

    def test_bullet(self):
        assert _parse_dot_status("●") == "N/A"


# ============================================================================
# Tests for build_heatmap_template
# ============================================================================
class TestBuildHeatmapTemplate:
    def test_returns_dataframe(self):
        df = build_heatmap_template()
        assert isinstance(df, pd.DataFrame)
        assert "Op Code" in df.columns
        assert "Operation Mode" in df.columns
        assert len(df) > 0

    def test_contains_known_operations(self):
        df = build_heatmap_template()
        op_codes = df["Op Code"].tolist()
        assert 10000000 in op_codes  # AVL-DRIVE Rating
        assert 10100000 in op_codes  # Drive away
        assert 10120200 in op_codes  # Constant load


# ============================================================================
# Tests for parse_heatmap_data
# ============================================================================
class TestParseHeatmapData:
    def test_basic_parse(self):
        text = (
            "\tOperation Modes\tVehicle A\tVehicle B\n"
            "\tOperation Modes\tDR\tDR\n"
            "10000000\tAVL-DRIVE Rating\t8.3\t8.2\n"
            "10100000\tDrive away\t7.9\t7.4\n"
        )
        result = parse_heatmap_data(text)
        assert result is not None
        assert len(result["vehicle_names"]) == 2
        assert len(result["data"]) == 2

    def test_empty_input(self):
        assert parse_heatmap_data("") is None
        assert parse_heatmap_data("   ") is None

    def test_separator_columns(self):
        """Data Transfer Sheet format has empty separator columns between vehicles."""
        text = (
            "\tOperation Modes\tBYD Atto 3 BEV_NORMAL\t\tBYD_Dolphin Surf_BEV_NORMAL\n"
            "\tOperation Modes\tDR\t\tDR\n"
            "10000000\tAVL-DRIVE Rating\t8.3\t\t8.2\n"
            "10100000\tDrive away\t7.9\t\t7.4\n"
            "10120300\tLoad increase\t\t\t8.8\n"
        )
        result = parse_heatmap_data(text)
        assert result is not None
        assert result["vehicle_names"] == [
            "BYD Atto 3 BEV_NORMAL",
            "BYD_Dolphin Surf_BEV_NORMAL",
        ]
        df = result["data"]
        assert len(df) == 3
        # First row: both vehicles have scores
        assert df.iloc[0]["BYD Atto 3 BEV_NORMAL"] == 8.3
        assert df.iloc[0]["BYD_Dolphin Surf_BEV_NORMAL"] == 8.2
        # Third row: only second vehicle has a score
        assert pd.isna(df.iloc[2]["BYD Atto 3 BEV_NORMAL"])
        assert df.iloc[2]["BYD_Dolphin Surf_BEV_NORMAL"] == 8.8

    def test_multiple_separator_columns(self):
        """Three vehicles each separated by empty columns."""
        text = (
            "\tOperation Modes\tVehicle A\t\tVehicle B\t\tVehicle C\n"
            "\tOperation Modes\tDR\t\tDR\t\tDR\n"
            "10000000\tAVL-DRIVE Rating\t8.3\t\t8.2\t\t7.5\n"
        )
        result = parse_heatmap_data(text)
        assert result is not None
        assert len(result["vehicle_names"]) == 3
        assert result["data"].iloc[0]["Vehicle A"] == 8.3
        assert result["data"].iloc[0]["Vehicle B"] == 8.2
        assert result["data"].iloc[0]["Vehicle C"] == 7.5

    def test_no_header_rows(self):
        """Data rows only, no header or DR rows."""
        text = "10000000\tAVL-DRIVE Rating\t8.3\t8.2\n10100000\tDrive away\t7.9\t7.4\n"
        result = parse_heatmap_data(text)
        assert result is not None
        assert len(result["vehicle_names"]) == 2
        assert len(result["data"]) == 2


# ============================================================================
# Tests for evaluate_avl_status
# ============================================================================
class TestEvaluateAVLStatus:
    def test_basic_evaluation(self):
        sheet1_data = {
            "target_car": "Car A",
            "tested_car": "Car B",
            "sections": [],
            "operations": [
                {
                    "section": "Drive away",
                    "op_code": 10101100,
                    "operation": "DASS",
                    "driv_p1": "GREEN",
                    "driv_p2": "GREEN",
                    "driv_p3": "GREEN",
                    "driv_tested": 99.8,
                    "driv_target": 100,
                    "resp_p1": "GREEN",
                    "resp_p2": "GREEN",
                    "resp_p3": "GREEN",
                    "resp_tested": 98,
                    "resp_target": 100,
                },
            ],
        }

        heatmap_df = pd.DataFrame({
            "Op Code": [10101100],
            "Operation Mode": ["Standing start"],
            "Car B": [7.6],
        })

        result = evaluate_avl_status(sheet1_data, heatmap_df, "Car A", "Car B")
        assert len(result) == 1
        assert result.iloc[0]["Driv Status"] == "GREEN"
        assert result.iloc[0]["Final Status"] == "GREEN"


# ============================================================================
# Tests for build_overall_status
# ============================================================================
class TestBuildOverallStatus:
    def test_all_green(self):
        df = pd.DataFrame({
            "Op Code": [10101100, 10101100],
            "Operation": ["Op1", "Op1"],
            "Final Status": ["GREEN", "GREEN"],
        })
        result = build_overall_status(df)
        assert result.iloc[0]["Overall Status"] == "GREEN"

    def test_any_red(self):
        df = pd.DataFrame({
            "Op Code": [10101100, 10101100],
            "Operation": ["Op1", "Op1"],
            "Final Status": ["GREEN", "RED"],
        })
        result = build_overall_status(df)
        assert result.iloc[0]["Overall Status"] == "RED"

    def test_all_na(self):
        df = pd.DataFrame({
            "Op Code": [10101100],
            "Operation": ["Op1"],
            "Final Status": ["N/A"],
        })
        result = build_overall_status(df)
        assert result.iloc[0]["Overall Status"] == "N/A"


# ============================================================================
# Tests for calculate_group_status
# ============================================================================
class TestCalculateGroupStatus:
    def test_all_green(self):
        ops = [
            {"driv_p1": "GREEN"},
            {"driv_p1": "GREEN"},
            {"driv_p1": "GREEN"},
        ]
        assert calculate_group_status(ops, "driv_p1") == "OK"

    def test_any_red(self):
        ops = [
            {"driv_p1": "GREEN"},
            {"driv_p1": "RED"},
            {"driv_p1": "GREEN"},
        ]
        assert calculate_group_status(ops, "driv_p1") == "NOK"

    def test_many_yellow(self):
        ops = [
            {"driv_p1": "YELLOW"},
            {"driv_p1": "YELLOW"},
            {"driv_p1": "GREEN"},
        ]
        # 66% yellow > 35% threshold
        assert calculate_group_status(ops, "driv_p1") == "Acceptable"

    def test_empty(self):
        assert calculate_group_status([], "driv_p1") == ""


# ============================================================================
# Tests for update_sub_operation_heatmap (parent group status + sub-op dots)
# ============================================================================
class TestUpdateSubOperationHeatmap:
    def _make_heatmap_df(self):
        """Build a small heatmap with one parent and two children."""
        return pd.DataFrame({
            "Op Code": [10100000, 10101300, 10101100, 10102400],
            "Operation Mode": ["Drive away", "Creep", "Standing start", "Rolling start"],
        })

    def _make_eval_results(self, statuses):
        """Build eval results for the three child op codes."""
        codes = [10101300, 10101100, 10102400]
        rows = []
        for code, status in zip(codes, statuses):
            rows.append({"Op Code": code, "Operation": "x", "Final Status": status})
        return pd.DataFrame(rows)

    def test_parent_ok_when_all_children_green(self):
        hm = self._make_heatmap_df()
        ev = self._make_eval_results(["GREEN", "GREEN", "GREEN"])
        result = update_sub_operation_heatmap(hm, ev)
        assert result.iloc[0]["Status"] == "OK"

    def test_parent_nok_when_any_child_red(self):
        hm = self._make_heatmap_df()
        ev = self._make_eval_results(["GREEN", "RED", "GREEN"])
        result = update_sub_operation_heatmap(hm, ev)
        assert result.iloc[0]["Status"] == "NOK"

    def test_parent_acceptable_when_many_yellow(self):
        hm = self._make_heatmap_df()
        ev = self._make_eval_results(["YELLOW", "YELLOW", "GREEN"])
        result = update_sub_operation_heatmap(hm, ev)
        # 2/3 = 66% > 35% -> Acceptable
        assert result.iloc[0]["Status"] == "Acceptable"

    def test_sub_operation_gets_status(self):
        hm = self._make_heatmap_df()
        ev = self._make_eval_results(["GREEN", "RED", "YELLOW"])
        result = update_sub_operation_heatmap(hm, ev)
        assert result.iloc[1]["Status"] == "GREEN"    # 10101300
        assert result.iloc[2]["Status"] == "RED"      # 10101100
        assert result.iloc[3]["Status"] == "YELLOW"   # 10102400

    def test_empty_eval_results(self):
        hm = self._make_heatmap_df()
        ev = pd.DataFrame()
        result = update_sub_operation_heatmap(hm, ev)
        assert all(result["Status"] == "")

    def test_parent_empty_when_no_children_have_status(self):
        hm = self._make_heatmap_df()
        ev = self._make_eval_results(["N/A", "N/A", "N/A"])
        result = update_sub_operation_heatmap(hm, ev)
        # Parent should still be empty when all children are N/A
        assert result.iloc[0]["Status"] == ""


# ============================================================================
# Tests for filter_heatmap_rows
# ============================================================================
class TestFilterHeatmapRows:
    def test_filter_removes_empty(self):
        df = pd.DataFrame({
            "Op Code": [1, 2, 3],
            "Operation Mode": ["A", "B", "C"],
            "Vehicle": [8.0, None, 7.5],
        })
        result = filter_heatmap_rows(df, "Vehicle")
        assert len(result) == 2
        assert 2 not in result["Op Code"].values


# ============================================================================
# Tests for parse_sheet1_from_excel
# ============================================================================
EXCEL_FILE = os.path.join(
    os.path.dirname(os.path.dirname(os.path.abspath(__file__))),
    "AVLDrive_Heatmap_Tool version_5.1_Atto3.xlsm",
)


class TestParseSheet1FromExcel:
    """Integration tests that read the actual sample Excel workbook."""

    def test_returns_dict_with_expected_keys(self):
        if not os.path.exists(EXCEL_FILE):
            return  # skip when the sample file isn't present
        result = parse_sheet1_from_excel(EXCEL_FILE)
        assert result is not None
        assert "target_car" in result
        assert "tested_car" in result
        assert "sections" in result
        assert "operations" in result

    def test_detects_car_names(self):
        if not os.path.exists(EXCEL_FILE):
            return
        result = parse_sheet1_from_excel(EXCEL_FILE)
        # The sample file has two vehicles in the header
        assert result["tested_car"] != ""
        assert result["target_car"] != ""
        assert result["tested_car"] != result["target_car"]

    def test_reads_green_dot(self):
        if not os.path.exists(EXCEL_FILE):
            return
        result = parse_sheet1_from_excel(EXCEL_FILE)
        # DA Rolling Start (10102400) has GREEN dots
        op = next((o for o in result["operations"] if o["op_code"] == 10102400), None)
        assert op is not None
        assert op["driv_p1"] == "GREEN"

    def test_reads_red_dot(self):
        if not os.path.exists(EXCEL_FILE):
            return
        result = parse_sheet1_from_excel(EXCEL_FILE)
        # Drive Away Creep (10101300) has RED driv_p1
        op = next((o for o in result["operations"] if o["op_code"] == 10101300), None)
        assert op is not None
        assert op["driv_p1"] == "RED"

    def test_reads_yellow_dot(self):
        if not os.path.exists(EXCEL_FILE):
            return
        result = parse_sheet1_from_excel(EXCEL_FILE)
        # One of the Maneuvering rows (10097800) has YELLOW resp_p1
        maneuvering_ops = [o for o in result["operations"] if o["op_code"] == 10097800]
        yellow_found = any(o["resp_p1"] == "YELLOW" for o in maneuvering_ops)
        assert yellow_found, "Expected at least one Maneuvering op with YELLOW resp_p1"

    def test_white_dot_is_na(self):
        if not os.path.exists(EXCEL_FILE):
            return
        result = parse_sheet1_from_excel(EXCEL_FILE)
        # Accel Cst Load (10120200) has white resp dots -> N/A
        op = next((o for o in result["operations"] if o["op_code"] == 10120200), None)
        assert op is not None
        assert op["resp_p1"] == "N/A"

    def test_sections_parsed(self):
        if not os.path.exists(EXCEL_FILE):
            return
        result = parse_sheet1_from_excel(EXCEL_FILE)
        section_names = [s["name"] for s in result["sections"]]
        assert len(section_names) > 0
        assert "Drive away" in section_names

    def test_missing_sheet_returns_none(self):
        """A workbook without Sheet1 returns None."""
        import tempfile
        import openpyxl
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "OtherSheet"
        ws["A1"] = "dummy"
        tmp = tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False)
        wb.save(tmp.name)
        wb.close()
        tmp.close()
        result = parse_sheet1_from_excel(tmp.name)
        assert result is None
        os.unlink(tmp.name)


# ============================================================================
# Tests for parse_odriv_from_excel
# ============================================================================
ODRIV_FILE = os.path.join(
    os.path.dirname(os.path.dirname(os.path.abspath(__file__))),
    "sample_data",
    "odriv_sample.xlsm",
)


class TestParseOdrivFromExcel:
    """Integration tests that read the ODRIV sample Excel workbook."""

    def test_returns_dict_with_expected_keys(self):
        if not os.path.exists(ODRIV_FILE):
            return
        from evaluation_engine import parse_odriv_from_excel
        result = parse_odriv_from_excel(ODRIV_FILE)
        assert result is not None
        assert "target_car" in result
        assert "tested_car" in result
        assert "sections" in result
        assert "operations" in result

    def test_detects_car_names(self):
        if not os.path.exists(ODRIV_FILE):
            return
        from evaluation_engine import parse_odriv_from_excel
        result = parse_odriv_from_excel(ODRIV_FILE)
        assert "BYD_Dolphin" in result["tested_car"]
        assert "Atto 3" in result["target_car"]

    def test_reads_green_dot(self):
        if not os.path.exists(ODRIV_FILE):
            return
        from evaluation_engine import parse_odriv_from_excel
        result = parse_odriv_from_excel(ODRIV_FILE)
        op = next((o for o in result["operations"] if o["op_code"] == 10102400), None)
        assert op is not None
        assert op["driv_p1"] == "GREEN"

    def test_reads_red_dot(self):
        if not os.path.exists(ODRIV_FILE):
            return
        from evaluation_engine import parse_odriv_from_excel
        result = parse_odriv_from_excel(ODRIV_FILE)
        op = next((o for o in result["operations"] if o["op_code"] == 10101300), None)
        assert op is not None
        assert op["driv_p1"] == "RED"

    def test_reads_yellow_dot(self):
        if not os.path.exists(ODRIV_FILE):
            return
        from evaluation_engine import parse_odriv_from_excel
        result = parse_odriv_from_excel(ODRIV_FILE)
        op = next((o for o in result["operations"] if o["op_code"] == 10451400), None)
        assert op is not None
        assert op["driv_p1"] == "YELLOW"

    def test_reads_responsiveness_dots(self):
        if not os.path.exists(ODRIV_FILE):
            return
        from evaluation_engine import parse_odriv_from_excel
        result = parse_odriv_from_excel(ODRIV_FILE)
        # Maneuvering (10097800) has YELLOW resp_p1
        maneuvering_ops = [o for o in result["operations"] if o["op_code"] == 10097800]
        yellow_found = any(o["resp_p1"] == "YELLOW" for o in maneuvering_ops)
        assert yellow_found

    def test_white_dot_is_na(self):
        if not os.path.exists(ODRIV_FILE):
            return
        from evaluation_engine import parse_odriv_from_excel
        result = parse_odriv_from_excel(ODRIV_FILE)
        # Accel Cst Load (10120200) has white resp dots -> N/A
        op = next((o for o in result["operations"] if o["op_code"] == 10120200), None)
        assert op is not None
        assert op["resp_p1"] == "N/A"

    def test_sections_parsed(self):
        if not os.path.exists(ODRIV_FILE):
            return
        from evaluation_engine import parse_odriv_from_excel
        result = parse_odriv_from_excel(ODRIV_FILE)
        section_names = [s["name"] for s in result["sections"]]
        assert len(section_names) >= 5
        assert "Drive away" in section_names

    def test_percentages_parsed(self):
        if not os.path.exists(ODRIV_FILE):
            return
        from evaluation_engine import parse_odriv_from_excel
        result = parse_odriv_from_excel(ODRIV_FILE)
        op = next((o for o in result["operations"] if o["op_code"] == 10101300), None)
        assert op is not None
        assert op["driv_tested"] > 0
        assert op["driv_target"] > 0

    def test_missing_rating_sheet_returns_none(self):
        """A workbook without RATING returns None."""
        import tempfile
        import openpyxl
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "OtherSheet"
        ws["A1"] = "dummy"
        tmp = tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False)
        wb.save(tmp.name)
        wb.close()
        tmp.close()
        from evaluation_engine import parse_odriv_from_excel
        result = parse_odriv_from_excel(tmp.name)
        assert result is None
        os.unlink(tmp.name)


# ============================================================================
# Tests for _normalize_for_match
# ============================================================================
class TestNormalizeForMatch:
    def test_basic_normalization(self):
        from evaluation_engine import _normalize_for_match
        assert _normalize_for_match("Drive Away") == "driveaway"
        assert _normalize_for_match("Drive-Away") == "driveaway"
        assert _normalize_for_match("Drive_Away") == "driveaway"

    def test_empty_string(self):
        from evaluation_engine import _normalize_for_match
        assert _normalize_for_match("") == ""


# ============================================================================
# Tests for _match_sheet_to_operation
# ============================================================================
class TestMatchSheetToOperation:
    def test_exact_match_section_hyphen_op(self):
        from evaluation_engine import _match_sheet_to_operation
        sheets = ["Driveaway-Creep", "Acceleration-ConstLoad", "TipIn"]
        assert _match_sheet_to_operation(sheets, "Drive away", "Creep") == "Driveaway-Creep"

    def test_section_underscore_op(self):
        from evaluation_engine import _match_sheet_to_operation
        sheets = ["Drive_away_Creep", "Other"]
        assert _match_sheet_to_operation(sheets, "Drive away", "Creep") == "Drive_away_Creep"

    def test_op_name_only_fallback(self):
        from evaluation_engine import _match_sheet_to_operation
        sheets = ["Creep", "Standing start"]
        assert _match_sheet_to_operation(sheets, "Drive away", "Creep") == "Creep"

    def test_no_match_returns_none(self):
        from evaluation_engine import _match_sheet_to_operation
        sheets = ["Acceleration", "TipIn"]
        assert _match_sheet_to_operation(sheets, "Drive away", "Creep") is None

    def test_empty_sheets_list(self):
        from evaluation_engine import _match_sheet_to_operation
        assert _match_sheet_to_operation([], "Drive away", "Creep") is None

    def test_fuzzy_contains_section_and_op(self):
        from evaluation_engine import _match_sheet_to_operation
        sheets = ["DA_Driveaway_Creep_Details"]
        assert _match_sheet_to_operation(sheets, "Drive away", "Creep") == "DA_Driveaway_Creep_Details"


# ============================================================================
# Tests for _extract_file_prefix
# ============================================================================
class TestExtractFilePrefix:
    def test_typical_avl_filename(self):
        from evaluation_engine import _extract_file_prefix
        result = _extract_file_prefix(
            "GS-RL_0%_Normal_Standard_BYD_Dolphin_Surf_BEV001_inca"
        )
        assert result == "GS_RL_0%"

    def test_cold_stop_word(self):
        from evaluation_engine import _extract_file_prefix
        result = _extract_file_prefix("DASS_Eng_On_Cold_BYD_Atto3")
        assert result == "DASS_Eng_On"

    def test_empty_string(self):
        from evaluation_engine import _extract_file_prefix
        assert _extract_file_prefix("") == ""

    def test_none_input(self):
        from evaluation_engine import _extract_file_prefix
        assert _extract_file_prefix(None) == ""

    def test_no_stop_word(self):
        from evaluation_engine import _extract_file_prefix
        result = _extract_file_prefix("MyTest_123_456")
        assert result == "MyTest_123_456"


# ============================================================================
# Tests for generate_red_comments
# ============================================================================
class TestGenerateRedComments:
    def test_generates_comment_for_red_p1_driv(self):
        from evaluation_engine import generate_red_comments

        sheet1_data = {
            "operations": [
                {
                    "op_code": 10101300,
                    "operation": "Creep",
                    "section": "Drive away",
                    "driv_p1": "RED",
                    "resp_p1": "N/A",
                },
            ],
        }

        odriv_details = {
            "Driveaway-Creep": [
                {
                    "file": "GS-RL_0%_Normal_Standard_BYD_Dolphin",
                    "criteria": "Brake Release Bump",
                    "priority": 1,
                    "rating": "Red",
                    "value": 6.2,
                },
                {
                    "file": "DASS_Normal_Standard_BYD_Dolphin",
                    "criteria": "Jerk Peak",
                    "priority": 1,
                    "rating": "Red+",
                    "value": 7.1,
                },
                {
                    "file": "DASS_Normal_Standard_BYD_Dolphin",
                    "criteria": "Smoothness",
                    "priority": 2,
                    "rating": "Red",
                    "value": 5.0,
                },
            ],
        }

        heatmap_df = pd.DataFrame({"Op Code": [10101300], "Status": ["RED"]})
        result = generate_red_comments(sheet1_data, heatmap_df, odriv_details)

        assert 10101300 in result
        comment = result[10101300]
        assert "Red P1 Drivability" in comment
        assert "Brake Release Bump" in comment
        assert "GS_RL_0%" in comment

    def test_no_comment_for_green_operation(self):
        from evaluation_engine import generate_red_comments

        sheet1_data = {
            "operations": [
                {
                    "op_code": 10101300,
                    "operation": "Creep",
                    "section": "Drive away",
                    "driv_p1": "GREEN",
                    "resp_p1": "GREEN",
                },
            ],
        }

        odriv_details = {"Driveaway-Creep": []}
        heatmap_df = pd.DataFrame({"Op Code": [10101300], "Status": ["GREEN"]})
        result = generate_red_comments(sheet1_data, heatmap_df, odriv_details)
        assert 10101300 not in result

    def test_no_matching_sheet_still_records_reason(self):
        from evaluation_engine import generate_red_comments

        sheet1_data = {
            "operations": [
                {
                    "op_code": 10101300,
                    "operation": "Creep",
                    "section": "Drive away",
                    "driv_p1": "RED",
                    "resp_p1": "N/A",
                },
            ],
        }

        odriv_details = {"SomeOtherSheet": []}
        heatmap_df = pd.DataFrame({"Op Code": [10101300], "Status": ["RED"]})
        result = generate_red_comments(sheet1_data, heatmap_df, odriv_details)
        assert 10101300 in result
        assert "Red P1 Drivability" in result[10101300]

    def test_empty_odriv_details(self):
        from evaluation_engine import generate_red_comments
        result = generate_red_comments({"operations": []}, pd.DataFrame(), {})
        assert result == {}

    def test_none_inputs(self):
        from evaluation_engine import generate_red_comments
        assert generate_red_comments(None, pd.DataFrame(), {}) == {}
        assert generate_red_comments({}, pd.DataFrame(), None) == {}

    def test_picks_lowest_value_event(self):
        from evaluation_engine import generate_red_comments

        sheet1_data = {
            "operations": [
                {
                    "op_code": 10101300,
                    "operation": "Creep",
                    "section": "Drive away",
                    "driv_p1": "RED",
                    "resp_p1": "N/A",
                },
            ],
        }

        odriv_details = {
            "Driveaway-Creep": [
                {"file": "File_A_Normal_X", "criteria": "CriteriaHigh", "priority": 1, "rating": "Red", "value": 8.0},
                {"file": "File_B_Normal_X", "criteria": "CriteriaLow", "priority": 1, "rating": "Red", "value": 5.5},
                {"file": "File_C_Normal_X", "criteria": "CriteriaMed", "priority": 1, "rating": "Red+", "value": 6.0},
            ],
        }

        heatmap_df = pd.DataFrame({"Op Code": [10101300], "Status": ["RED"]})
        result = generate_red_comments(sheet1_data, heatmap_df, odriv_details)
        assert "CriteriaLow" in result[10101300]

    def test_both_driv_and_resp_red(self):
        from evaluation_engine import generate_red_comments

        sheet1_data = {
            "operations": [
                {
                    "op_code": 10101300,
                    "operation": "Creep",
                    "section": "Drive away",
                    "driv_p1": "RED",
                    "resp_p1": "RED",
                },
            ],
        }

        odriv_details = {
            "Driveaway-Creep": [
                {"file": "F_Normal_X", "criteria": "C1", "priority": 1, "rating": "Red", "value": 6.0},
            ],
            "Driveaway-Creep__resp": [
                {"file": "F_Normal_Y", "criteria": "Response delay", "priority": 1, "rating": "Red", "value": 5.5},
            ],
        }

        heatmap_df = pd.DataFrame({"Op Code": [10101300], "Status": ["RED"]})
        result = generate_red_comments(sheet1_data, heatmap_df, odriv_details)
        comment = result[10101300]
        assert "Red P1 Drivability" in comment
        assert "C1" in comment
        assert "Red P1 Responsiveness" in comment
        assert "Response delay" in comment

    def test_resp_only_red_uses_resp_section(self):
        """When only Responsiveness is RED, use the __resp detail data."""
        from evaluation_engine import generate_red_comments

        sheet1_data = {
            "operations": [
                {
                    "op_code": 10101300,
                    "operation": "Creep",
                    "section": "Drive away",
                    "driv_p1": "GREEN",
                    "resp_p1": "RED",
                },
            ],
        }

        odriv_details = {
            "Driveaway-Creep": [
                {"file": "F_Driv", "criteria": "Bump", "priority": 1, "rating": "Red", "value": 6.0},
            ],
            "Driveaway-Creep__resp": [
                {"file": "F_Resp", "criteria": "Response delay", "priority": 1, "rating": "Red", "value": 5.0},
            ],
        }

        heatmap_df = pd.DataFrame({"Op Code": [10101300], "Status": ["RED"]})
        result = generate_red_comments(sheet1_data, heatmap_df, odriv_details)
        comment = result[10101300]
        assert "Red P1 Responsiveness" in comment
        assert "Response delay" in comment
        # Should NOT include Drivability info
        assert "Red P1 Drivability" not in comment
        assert "Bump" not in comment

    def test_avl_below_7_adds_comment(self):
        """When AVL < 7 causes RED status, 'AVL<7' should appear in comment."""
        from evaluation_engine import generate_red_comments

        sheet1_data = {
            "operations": [
                {
                    "op_code": 10101300,
                    "operation": "Drive away",
                    "section": "Drive away",
                    "driv_p1": "GREEN",
                    "resp_p1": "GREEN",
                },
            ]
        }
        heatmap_df = pd.DataFrame()
        eval_results_df = pd.DataFrame([
            {"Op Code": 10101300, "Tested AVL": 5.5},
        ])
        result = generate_red_comments(
            sheet1_data, heatmap_df, None,
            eval_results_df=eval_results_df,
        )
        assert 10101300 in result
        assert "AVL<7" in result[10101300]

    def test_avl_below_7_with_red_p1_combines_reasons(self):
        """When both AVL<7 and P1=RED, comment should include both reasons."""
        from evaluation_engine import generate_red_comments

        sheet1_data = {
            "operations": [
                {
                    "op_code": 10101300,
                    "operation": "Drive away",
                    "section": "Drive away",
                    "driv_p1": "RED",
                    "resp_p1": "GREEN",
                },
            ]
        }
        heatmap_df = pd.DataFrame()
        odriv_details = {
            "Drive away": [
                {"file": "F_Normal_X", "criteria": "Bump", "priority": 1,
                 "rating": "Red", "value": 6.0},
            ]
        }
        eval_results_df = pd.DataFrame([
            {"Op Code": 10101300, "Tested AVL": 6.0},
        ])
        result = generate_red_comments(
            sheet1_data, heatmap_df, odriv_details,
            eval_results_df=eval_results_df,
        )
        comment = result[10101300]
        assert "AVL<7" in comment
        assert "Red P1 Drivability" in comment

    def test_avl_at_7_no_avl_comment(self):
        """When AVL == 7 (exactly at threshold), no 'AVL<7' comment."""
        from evaluation_engine import generate_red_comments

        sheet1_data = {
            "operations": [
                {
                    "op_code": 10101300,
                    "operation": "Drive away",
                    "section": "Drive away",
                    "driv_p1": "GREEN",
                    "resp_p1": "GREEN",
                },
            ]
        }
        eval_results_df = pd.DataFrame([
            {"Op Code": 10101300, "Tested AVL": 7.0},
        ])
        result = generate_red_comments(
            sheet1_data, pd.DataFrame(), None,
            eval_results_df=eval_results_df,
        )
        assert 10101300 not in result

    def test_avl_below_7_no_eval_results(self):
        """Without eval_results_df, AVL<7 comments are not generated."""
        from evaluation_engine import generate_red_comments

        sheet1_data = {
            "operations": [
                {
                    "op_code": 10101300,
                    "operation": "Drive away",
                    "section": "Drive away",
                    "driv_p1": "GREEN",
                    "resp_p1": "GREEN",
                },
            ]
        }
        result = generate_red_comments(
            sheet1_data, pd.DataFrame(), None,
        )
        assert 10101300 not in result


# ============================================================================
# Tests for _parse_detail_sheet header disambiguation
# ============================================================================
class TestParseDetailSheetHeaders:
    """Ensure 'Event Rating' column is not stolen by the 'event' keyword in criteria."""

    def test_event_rating_not_captured_as_criteria(self):
        import openpyxl
        from evaluation_engine import _parse_detail_sheet

        wb = openpyxl.Workbook()
        ws = wb.active
        # Header row: File | Event | Priority | Event Rating | DR
        ws.cell(row=1, column=1, value="File")
        ws.cell(row=1, column=2, value="Event")
        ws.cell(row=1, column=3, value="Priority")
        ws.cell(row=1, column=4, value="Event Rating")
        ws.cell(row=1, column=5, value="DR")
        # Data row
        ws.cell(row=2, column=1, value="GS-RL_0%_Normal_BYD")
        ws.cell(row=2, column=2, value="Brake Release Bump")
        ws.cell(row=2, column=3, value=1)
        ws.cell(row=2, column=4, value="Red")
        ws.cell(row=2, column=5, value=6.2)

        events = _parse_detail_sheet(ws)
        wb.close()

        assert len(events) == 1
        e = events[0]
        assert e["criteria"] == "Brake Release Bump"
        assert e["rating"] == "Red"
        assert e["value"] == 6.2
        assert e["priority"] == 1
        assert "GS-RL_0%" in e["file"]


class TestParseDetailSheetWideFormat:
    """Tests for ODRIV wide-format detail sheets where criteria are column headers."""

    def test_wide_format_finds_worst_criteria(self):
        """Simulates ODRIV wide-format: criteria names as column headers."""
        import openpyxl
        from openpyxl.styles import PatternFill
        from evaluation_engine import _parse_detail_sheet

        green_fill = PatternFill(patternType="solid", fgColor="FF009926")
        red_fill = PatternFill(patternType="solid", fgColor="FFF66E60")

        wb = openpyxl.Workbook()
        ws = wb.active
        # Header row (mimicking ODRIV structure)
        ws.cell(row=1, column=13, value="Criticality")
        ws.cell(row=1, column=14, value="Event Priority")
        ws.cell(row=1, column=15, value="Event Rating")
        ws.cell(row=1, column=16, value="Throttle Position")
        ws.cell(row=1, column=20, value="Brake release delay")
        ws.cell(row=1, column=21, value="Brake release bump")
        ws.cell(row=1, column=22, value="Acceleration disturbances")
        ws.cell(row=1, column=23, value="Start time of the Sub Event")
        ws.cell(row=1, column=24, value="Acquisition Name")

        # Data row: RED P1 event with worst score in "Brake release bump"
        ws.cell(row=2, column=13, value=1)
        ws.cell(row=2, column=14, value=1)
        ws.cell(row=2, column=15, value="RED")
        ws.cell(row=2, column=16, value=0)        # Throttle: no fill
        c = ws.cell(row=2, column=20, value=8.3);  c.fill = green_fill
        c = ws.cell(row=2, column=21, value=6.4);  c.fill = red_fill   # worst
        c = ws.cell(row=2, column=22, value=9.3);  c.fill = green_fill
        ws.cell(row=2, column=23, value=5.39)      # Start time: no fill
        ws.cell(row=2, column=24, value="DriveAway_0%_Normal_Standard_BYD_Dolphin")

        events = _parse_detail_sheet(ws)
        wb.close()

        assert len(events) >= 1
        e = events[0]
        assert e["priority"] == 1
        assert e["rating"] == "RED"
        assert e["criteria"] == "Brake release bump"
        assert e["value"] == 6.4
        assert "DriveAway_0%" in e["file"]

    def test_wide_format_start_time_not_criteria(self):
        """Start time of the Sub Event should NOT be detected as criteria."""
        import openpyxl
        from openpyxl.styles import PatternFill
        from evaluation_engine import _parse_detail_sheet

        red_fill = PatternFill(patternType="solid", fgColor="FFF66E60")

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.cell(row=1, column=14, value="Event Priority")
        ws.cell(row=1, column=15, value="Event Rating")
        ws.cell(row=1, column=20, value="Brake release bump")
        ws.cell(row=1, column=21, value="Start time of the Sub Event")
        ws.cell(row=1, column=22, value="Acquisition Name")

        ws.cell(row=2, column=14, value=1)
        ws.cell(row=2, column=15, value="RED")
        c = ws.cell(row=2, column=20, value=6.0);  c.fill = red_fill
        ws.cell(row=2, column=21, value=1.5)  # lower, but no fill
        ws.cell(row=2, column=22, value="TestFile_Normal_BYD")

        events = _parse_detail_sheet(ws)
        wb.close()

        assert len(events) >= 1
        e = events[0]
        # Criteria should be "Brake release bump" not start time
        assert e["criteria"] == "Brake release bump"
        assert e["value"] == 6.0

    def test_wide_format_ignores_uncolored_lower_value(self):
        """Uncolored cells with lower values should be ignored."""
        import openpyxl
        from openpyxl.styles import PatternFill
        from evaluation_engine import _parse_detail_sheet

        green_fill = PatternFill(patternType="solid", fgColor="FF009926")
        yellow_fill = PatternFill(patternType="solid", fgColor="FFCFE747")

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.cell(row=1, column=14, value="Event Priority")
        ws.cell(row=1, column=15, value="Event Rating")
        ws.cell(row=1, column=16, value="Throttle Position")
        ws.cell(row=1, column=17, value="Front EM Speed")
        ws.cell(row=1, column=20, value="Brake release bump")
        ws.cell(row=1, column=21, value="Acceleration disturbances")
        ws.cell(row=1, column=22, value="Acquisition Name")

        ws.cell(row=2, column=14, value=1)
        ws.cell(row=2, column=15, value="RED")
        ws.cell(row=2, column=16, value=0)           # 0 is lower, but no fill
        ws.cell(row=2, column=17, value=0)            # 0 is lower, but no fill
        c = ws.cell(row=2, column=20, value=7.5);  c.fill = yellow_fill
        c = ws.cell(row=2, column=21, value=8.0);  c.fill = green_fill
        ws.cell(row=2, column=22, value="Test_File_BYD")

        events = _parse_detail_sheet(ws)
        wb.close()

        assert len(events) >= 1
        e = events[0]
        # Should pick "Brake release bump" (7.5) not Throttle/EM Speed (0)
        assert e["criteria"] == "Brake release bump"
        assert e["value"] == 7.5
